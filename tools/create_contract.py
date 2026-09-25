from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = '/Users/epaphradito/develop/web/digitofina-web/DigitoFina_Universal_Software_Services_Agreement.docx'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    border = tbl_pr.first_child_found_in('w:tblBorders')
    if border is None:
        border = OxmlElement('w:tblBorders')
        tbl_pr.append(border)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = border.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            border.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), 'D9D9D9')


def set_cell_text(cell, text, bold=False, color='000000'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
for style_name, size in [('Title', 22), ('Heading 1', 14), ('Heading 2', 11.5)]:
    style = styles[style_name]
    style.font.name = 'Aptos Display' if style_name == 'Title' else 'Aptos'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(size)
    style.font.bold = True

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Universal Software Services Agreement')
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run('DigitoFina platform licensing, hosting, support, and customization')
run.italic = True
run.font.size = Pt(10.5)

add_body(doc, 'This agreement is a reusable commercial template. Complete the fields in the Commercial Schedule before signature, and have the final version reviewed under applicable Ugandan law where appropriate.')

add_heading(doc, 'Parties and effective date')
table = doc.add_table(rows=5, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True
borders(table)
fields = [
    ('Provider', 'DigitoFina / [legal entity name]'),
    ('Provider address', '[address, email, telephone]'),
    ('Customer', '[customer legal name]'),
    ('Customer address', '[address, email, telephone]'),
    ('Effective date', '[day] [month] [year]'),
]
for row, (label, value) in zip(table.rows, fields):
    set_cell_text(row.cells[0], label, bold=True, color='1F4E79')
    shade(row.cells[0], 'EAF2F8')
    set_cell_text(row.cells[1], value)

add_heading(doc, 'Purpose and scope')
add_body(doc, 'The Provider will make the DigitoFina financial management platform available to the Customer for the subscription or license period stated in the Commercial Schedule. The platform may include organization management, branches, users, clients, savings, loans, accounting, reporting, integrations, and other enabled modules.')

add_heading(doc, 'Commercial terms')
add_body(doc, 'The Customer will pay the amounts and on the dates set out in the Commercial Schedule. The Provider may suspend non-essential services after giving written notice of overdue amounts and a reasonable opportunity to cure the default.')
add_body(doc, 'Annual hosting is a separate recurring charge of UGX 200,000 unless the Commercial Schedule states a different written amount. Hosting covers the production environment, routine infrastructure operation, backups, security maintenance, and ordinary platform availability. Third-party transaction, messaging, payment, domain, or gateway charges are paid by the Customer unless expressly included.')
add_body(doc, 'Where the parties agree to a discounted introductory fee, that discount is limited to the stated term and does not create a permanent price guarantee. At the end of the term, the parties will review pricing with reference to the number of branches, users, clients, enabled modules, transaction volume, storage, integrations, and support or customization requirements.')

add_heading(doc, 'Provider responsibilities')
for item in [
    'Provide access to the agreed DigitoFina modules and maintain the production service with reasonable care.',
    'Apply security updates, maintain routine backups, and use reasonable measures to protect Customer data.',
    'Respond to reported faults within a reasonable time and communicate material service incidents where practical.',
    'Provide ordinary user guidance and reasonable product updates during the term.',
    'Discuss requested customizations in good faith and confirm any separate scope, price, and delivery date before starting chargeable work.',
    'Keep Customer information confidential, subject to lawful disclosure and the data processing terms below.',
]:
    add_bullet(doc, item)

add_heading(doc, 'Customer responsibilities')
for item in [
    'Pay all agreed fees and provide accurate billing, organization, branch, user, and contact information.',
    'Use the platform lawfully and only for the Customer’s own authorized operations.',
    'Keep passwords, devices, API credentials, and administrator access secure, and promptly report suspected compromise.',
    'Ensure that Customer data is accurate, that users have appropriate permissions, and that the Customer has authority to process submitted personal or financial data.',
    'Provide timely decisions, test data, approvals, and access reasonably needed for agreed configuration or customization work.',
    'Maintain its own internet connectivity, devices, local policies, and any third-party services it chooses to connect to the platform.',
]:
    add_bullet(doc, item)

add_heading(doc, 'Data ownership and confidentiality')
add_body(doc, 'The Customer retains ownership of data it submits to the platform. The Provider may process that data only to operate, secure, support, improve, and lawfully administer the service. Each party will protect the other party’s confidential information using reasonable safeguards and will disclose it only to personnel, advisers, or service providers who need it and are bound by confidentiality duties.')
add_body(doc, 'The Provider may use aggregated and de-identified operational statistics for service monitoring and improvement, provided those statistics do not identify the Customer, its clients, or individual transactions.')

add_heading(doc, 'Support, updates, and customization')
add_body(doc, 'Routine updates, bug fixes, security improvements, and ordinary configuration are included unless the Commercial Schedule states otherwise. Custom features, data migration, special reports, integrations, training beyond the agreed scope, and onsite work may be quoted separately. A written change order should state the requested work, price, assumptions, delivery target, acceptance method, and any effect on support or hosting charges.')

add_heading(doc, 'Availability and service limits')
add_body(doc, 'The Provider will use reasonable efforts to keep the service available, but planned maintenance, emergency maintenance, internet failures, third-party outages, force majeure events, and Customer-controlled systems may affect availability. The platform is a business tool and does not replace the Customer’s legal, accounting, credit, data protection, or internal control responsibilities.')

add_heading(doc, 'Term, renewal, and review')
add_body(doc, 'This agreement begins on the Effective Date and continues for the period in the Commercial Schedule. The parties may renew it in writing. Unless otherwise stated, the Provider will give at least 30 days’ notice of proposed renewal pricing, and the parties will review the commercial terms before the next period begins.')

add_heading(doc, 'Suspension and termination')
add_body(doc, 'Either party may terminate for a material breach that remains uncured 30 days after written notice. The Provider may suspend access immediately where necessary to protect the platform, comply with law, address a security threat, or respond to serious misuse. On termination, the Customer will pay outstanding amounts, and the Provider will make a reasonable export of Customer data available for [30] days, subject to payment of agreed migration or export costs. After that period, the Provider may delete the data in accordance with its retention policy and applicable law.')

add_heading(doc, 'Liability')
add_body(doc, 'Each party remains responsible for its own acts and omissions. To the extent permitted by law, neither party is liable for indirect, incidental, special, or consequential loss. The Provider’s total liability arising from this agreement will not exceed the fees paid or payable by the Customer during the 12 months before the event giving rise to the claim, except for fraud, wilful misconduct, confidentiality breaches, or liability that cannot lawfully be limited.')

add_heading(doc, 'Governing law and notices')
add_body(doc, 'This agreement is governed by the laws of Uganda. The parties will first try to resolve disputes through good-faith senior management discussions. Formal notices may be sent to the addresses and email contacts in the Commercial Schedule and are effective when delivery is reasonably confirmed.')

add_heading(doc, 'Commercial Schedule')
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(table)
headers = table.rows[0].cells
set_cell_text(headers[0], 'Commercial item', bold=True, color='FFFFFF')
set_cell_text(headers[1], 'Agreed term', bold=True, color='FFFFFF')
shade(headers[0], '1F4E79')
shade(headers[1], '1F4E79')
items = [
    ('License or service period', '[24 months from the Effective Date]'),
    ('Total introductory license or service fee', '[UGX __________]'),
    ('Amount paid on signing', '[UGX __________]'),
    ('Balance and due date', '[UGX __________ due by __________]'),
    ('Annual hosting fee', 'UGX 200,000 per year unless changed here: [UGX __________]'),
    ('Included branches', '[__________]'),
    ('Included users', '[__________]'),
    ('Included clients or records', '[__________]'),
    ('Included modules', '[list modules or attach an approved scope]'),
    ('Pricing review date', '[__________]'),
    ('Special payment or discount terms', '[__________]'),
]
for label, value in items:
    row = table.add_row()
    set_cell_text(row.cells[0], label, bold=True)
    set_cell_text(row.cells[1], value)

add_heading(doc, 'Signatures')
add_body(doc, 'By signing below, each party confirms that it has authority to enter into this agreement and accepts the Commercial Schedule and the terms above.')
table = doc.add_table(rows=5, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(table)
signature_fields = [
    ('For the Provider', 'For the Customer'),
    ('Name: ______________________________', 'Name: ______________________________'),
    ('Title: _______________________________', 'Title: _______________________________'),
    ('Signature: ____________________________', 'Signature: ____________________________'),
    ('Date: _________________________________', 'Date: _________________________________'),
]
for i, (left, right) in enumerate(signature_fields):
    set_cell_text(table.rows[i].cells[0], left, bold=(i == 0), color='1F4E79' if i == 0 else '000000')
    set_cell_text(table.rows[i].cells[1], right, bold=(i == 0), color='1F4E79' if i == 0 else '000000')
    if i == 0:
        shade(table.rows[i].cells[0], 'EAF2F8')
        shade(table.rows[i].cells[1], 'EAF2F8')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('DigitoFina Universal Software Services Agreement  |  Template for completion before signature').font.size = Pt(8)

doc.save(OUT)
print(OUT)
